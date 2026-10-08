from pathlib import Path
import os
import json,csv,hashlib,shutil
from collections import Counter
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

DEST=Path(os.environ.get('W8_BUNDLE_DIR',Path(__file__).resolve().parents[1]))
DATA=json.loads((DEST/'w8_content.json').read_text());SL=DATA['slides']
sources=json.loads((DEST/'research/source_registry.json').read_text())
for r in sources:
 r['slides']=[s['id'] for s in SL if r['id'] in s['sources']]
 r['ppt_use']='화면·대본·교사용 상세자료' if r['slides'] else '자료 보존·PPT 미반영'
 if r['id']=='D14':r['rechecked_url']='https://www.law.go.kr/LSW//lsSideInfoP.do?docCls=jo&joBrNo=00&joNo=0025&lsiSeq=283839&urlMode=lsScJoRltInfoR'
 if r['id']=='D38':r['rechecked_url']='https://www.law.go.kr/LSW/lsLinkCommonInfo.do?chrClsCd=010202&lsJoLnkSeq=1031357295'
 if r['id'] in ['D09','D14','D15','D38','F01','F02','F03','F04','S13','X01','X02']:r['w8_rechecked']='2026-10-08'
(DEST/'research/source_registry.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2)+'\n')
REF={r['id']:r for r in sources}
with (DEST/'research/source_registry.csv').open('w',encoding='utf-8-sig',newline='') as f:
 fields=['id','collection','title','by','date','url','status','scope','finding','limit','slides','ppt_use','same_url_ids'];w=csv.DictWriter(f,fields,extrasaction='ignore');w.writeheader();w.writerows(sources)
def table(headers,rows):return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join('---' for _ in headers)+'|']+['| '+' | '.join(map(str,r))+' |' for r in rows])
CH=[]
for name in dict.fromkeys(s['chapter'] for s in SL[:44]):
 rows=[s for s in SL[:44] if s['chapter']==name]
 CH.append([name,f"{rows[0]['page']}–{rows[-1]['page']}",f"{sum(s['minutes'] for s in rows):g}분"])
purposes=['입력 대상과 검색 대상, 경보의 의미를 구분한다.','표준 기록·평가·관제의 발전과 실제 배치의 과제를 연결한다.','영상 처리와 정지·체포의 각각 요건을 이해한다.','운영 실적·오류 분모·추가효과·비용을 나누어 읽는다.','오류의 전파, 시민 부담, 목적·비례와 구제를 살핀다.','국내 의견과 해외 규범을 토대로 수단·책임을 구체화한다.','앞 자료를 논거로 회수하고 자기 결론과 변경 조건을 작성한다.']
for row,p in zip(CH,purposes):row.append(p)
out=['# 8주차 목차안','**공공 CCTV의 실시간 AI 분석**','범죄·공공안전세미나Ⅱ · 검토용 초안 v1 · 2026-10-08',f"**논제** {SL[0]['rows'][0][1]}",'## 구성과 수업 운영','본편 44장, 보충 9장(보충 구분장 포함)이다. 강사는 허용 여부의 결론을 대신 제시하지 않고, 기술·법·실제 운영·권리·통제 자료를 학생 판단으로 연결한다.','설명과 표 읽기·판단 활동 88분에 질의 7분을 더한 **95분 운영안**이다. 실제 낭독·리허설 측정값이 아니다. 긴급 질문이나 보충 자료 사용 시 아래 질의시간과 본편 확장 구간을 조정한다.',table(['단원','PPT 쪽','운영 초안','역할'],CH),'## 질의 7분','1. p.15 뒤 2분: 얼굴 대조와 정지·체포에서 각각 필요한 정보는 무엇인가?\n2. p.23 뒤 2분: 같은 오경보 10건을 어떤 분모로 볼 것인가?\n3. p.37 뒤 3분: 목적·명단·장소가 바뀔 때 누구의 재심사가 필요한가?','## 종합과 학생 판단','p.38–44의 7장은 먼저 예약한 종합 구간이다. 정책 선택, 필요성·성과, 법적 근거·현장 조치, 오류·기록, 감독·중단을 정리한 뒤 학생이 결론·핵심 근거·남는 부담·판단 변경 조건을 작성한다.','## 페이지별 목차',table(['쪽','배치','제목'],[[s['page'],s['placement'],s['title']] for s in SL]),'## 표준구조 8기능과의 연결',table(['기능','반영 쪽'],[['논제와 출발점','1–2'],['범위와 핵심용어','3–6'],['현행 법·제도의 기준선','11–15'],['핵심 쟁점 지도','39–43에서 자료를 종합해 회수'],['정책대안·비교기준','21, 35–38'],['쟁점별 자료·논거','7–34, 39–42'],['종합·남는 판단','38–43'],['학생의 자기 판단','44']]),'## 자료와 제작 기준','기준 Git은 master `7b9d65fa7bcb27add2a53d04f2ea5b769d3310c9`이다. 기존 56건, 역사·사례 보완 22건, 제작 중 추가 법령 2건을 합한 80개 자료항목을 보존했다. 같은 사건을 다룬 서로 다른 문서는 함께 남기고 사용 관계를 대장에 표시한다. 80개를 모두 본편에 넣지는 않았다.','기존 W2·W3·W4는 분량을 비교하는 참고본으로 읽었고, W2 실제 PPT의 표지·목차·행·비교·종합 화면을 렌더링해 서식을 대조했다. 본편 장수는 자료의 역할에 맞춰 결정했다.','학생 화면에는 상세 서지·URL·검수 메모를 넣지 않았다. PPT 노트는 낭독 대본 / 출처 / 강사용 상세자료 / 내부 제작 메모로 구성했다.','국내 실시간 1:N 수사에 대한 경보–접촉–체포–기소–유죄 연결 원표와 직접 인과효과는 미확인이다. D17 고시 전문의 접근 실패를 유지하고 미확인 조문 내용을 인용하지 않았다.','## 함께 읽을 파일','- `outputs/W8_AI_CCTV_draft_v1.pptx`: 학생 화면과 발표자 노트\n- `outputs/W8_outline_script_draft_v1.docx`: 목차·페이지별 낭독 대본·추가 질문\n- `W8_page_plan.md`: 화면 문안·대본·출처·강사용 상세자료\n- `research/source_registry.csv`: 모든 수집자료와 PPT 연결\n- `materials/`: 기존 조사본과 역사·사례 보완본 원본\n- `qa/`: 페이지별 내용·렌더링 검수 기록']
(DEST/'W8_outline.md').write_text('\n\n'.join(out)+'\n')

plan=['# 8주차 페이지별 PPT 구성안과 대본','검토용 초안 v1 · 2026-10-08 · 본편 44장 + 보충 9장','화면 문안·낭독 대본과 강사용 상세자료·내부 제작 메모를 분리했다. 출처 ID는 통합 대장에 연결된다.']
script=['# 8주차 낭독 대본','검토용 초안 v1 · PPT 쪽수와 동일한 순서']
for s in SL:
 plan.extend([f"## {s['page']:02} {s['title']}",f"**단원·배치** {s['chapter']} / {s['placement']}\n\n**기억목표** {s['goal']}\n\n**학생 질문** {s['question']}",f"### 화면\n\n{s['band']}"])
 if s['headers']:plan.append(table(s['headers'],s['rows']))
 else:plan.extend([' — '.join(r) for r in s['rows']])
 plan.extend(s['below'])
 plan.append('### 낭독 대본\n\n'+'\n\n'.join(s['script']))
 plan.append('### 강사용 상세자료')
 for id in s['sources']:
  r=REF[id];plan.append(f"**{id} {r['by']} · {r['date']}**\n\n[{r['title']}]({r.get('rechecked_url',r['url'])})\n\n확인내용: {r['finding']}\n\n읽은 범위: {r['scope']}\n\n적용 범위·한계: {r.get('limit','')}")
 plan.append(f"### 내부 제작 메모\n\n논제 관계: {s['role']}\n\n{s['memo']}\n\n이전 연결: {s['previous_connection']}\n\n다음 연결: {s['next_connection']}")
 script.extend([f"## {s['page']:02} {s['title']}",'\n\n'.join(s['script']),f"추가 질문: {s['question']}",f"다음 연결: {s['next_connection']}",'출처: '+', '.join(f"[{id}]({REF[id]['url']})" for id in s['sources'])])
(DEST/'W8_page_plan.md').write_text('\n\n'.join(plan)+'\n')
(DEST/'W8_script.md').write_text('\n\n'.join(script)+'\n')

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.left_margin=sec.right_margin=Inches(.70);sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.60)
sec.footer_distance=Inches(.28)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[n];st.font.name='Pretendard';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Pretendard')
 st.paragraph_format.widow_control=True
 st.font.italic=False;st.font.underline=False
 if n in ['Normal','Subtitle','Caption']:st.font.bold=False
for st in doc.styles:
 for border in st._element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.18
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
for n,size in [('Title',27),('Subtitle',14),('Heading 1',18),('Heading 2',13),('Caption',10)]:doc.styles[n].font.size=Pt(size)
for n in ['Heading 1','Heading 2']:doc.styles[n].paragraph_format.space_before=Pt(10);doc.styles[n].paragraph_format.space_after=Pt(6)
doc.core_properties.title='8주차 공공 CCTV 실시간 AI 분석 목차와 대본안';doc.core_properties.author=''
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('8주차 검토용 초안  ·  ').font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
def para(t,style=None):return doc.add_paragraph(t,style)
def hyperlink(p,text,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Pretendard');font.set(qn('w:eastAsia'),'Pretendard');pr.append(font);color=OxmlElement('w:color');color.set(qn('w:val'),'1E2761');pr.append(color);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
def doc_table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for vals in [headers]+rows:
  cells=t.rows[0].cells if vals is headers else t.add_row().cells
  for j,v in enumerate(vals):
   c=cells[j];c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);m=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');m.append(e)
   pr.append(m)
   if vals is headers:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F4F8');pr.append(sh)
   c.text=str(v)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
    for run in p.runs:run.font.size=Pt(10.5);run.bold=vals is headers
 return t

para('8주차 공공 CCTV\n실시간 AI 분석','Title')
para('목차와 낭독 대본안','Subtitle')
para('범죄·공공안전세미나Ⅱ  /  검토용 초안 v1  /  2026년 10월 8일','Caption')
para(SL[0]['rows'][0][1])
para('기술의 작동과 국내 법적 기준, 실제 현장 성과와 권리 부담을 학습한 뒤 허용 범위와 통제 조건을 판단하는 강사용 초안이다. 강사는 자료와 서로 다른 논거를 제시하고, 학생은 자신의 결론과 그 이유를 구성한다.')
para('목차','Heading 1')
doc_table(['단원','PPT','운영안'],[r[:3] for r in CH],[4.8,1.0,1.0])
para('본편 44장과 보충 9장으로 구성했다. 보충 자료는 기술·사례·국내 제도에 관한 질문이 있을 때 선택한다.','Caption')
doc.add_page_break()
para('대본 사용과 시간 운영','Heading 1')
para('각 항목의 번호는 PPT 쪽수와 일치한다. 낭독 문장 다음의 추가 질문과 연결 문장은 수업 진행용이다. 문장 전체를 계속 읽기보다 화면의 표·절차를 함께 살펴보고 질문에 따라 설명을 확장한다.')
para('시간은 설명과 표 읽기·판단 활동 88분, 질의 7분을 합한 95분 운영안이다. 실제 리허설 측정값은 아니다. 학생 발표나 토론 시간이 포함되는 운영에서는 본편 확장 구간과 보충 자료를 조정한다.')
doc_table(['질의 시점','시간','질문'],[['p.15 뒤','2분','얼굴 대조와 정지·체포에서 각각 필요한 정보는 무엇인가?'],['p.23 뒤','2분','같은 오경보 10건을 어떤 분모로 볼 것인가?'],['p.37 뒤','3분','목적·명단·장소가 바뀔 때 누구의 재심사가 필요한가?']],[1.0,.65,5.15])
para('종합과 학생 판단','Heading 2')
para('p.38–44에서 앞서 본 자료를 회수한다. 학생은 자기 결론을 좌우하는 질문을 고르고, 핵심 자료 2개와 그 자료에서 결론으로 이어지는 이유를 적는다. 자신의 선택이 남기는 부담과 판단을 바꿀 조건도 함께 제시한다.')
para('상세자료와 출처','Heading 2')
para('각 대본의 출처 ID를 누르면 원자료가 열린다. 문서의 성격과 실제로 읽은 범위, 적용 한계는 W8_page_plan.md와 PPT 발표자 노트에 기록했다. 통합 대장은 기존 조사 56건, 역사·사례 보완 22건, 추가 법령 2건 등 80개 자료항목을 보존한다.')
para('국내 실시간 1:N 수사에 관한 경보부터 유죄까지의 연결 원표와 직접 인과효과는 미확인이다. 행정안전부 관련 고시 D17은 전문 접근 실패를 남겼고, 확인하지 못한 보존기간 등 세부 요건을 인용하지 않았다.')
para('기준 저장소','Heading 2')
para('jskwon89/lecture  ·  master 7b9d65f  ·  2026년 10월 8일 08:41 KST','Caption')
for i,s in enumerate(SL):
 if i%3==0:doc.add_page_break()
 para(f"{s['page']:02}  {s['title']}",'Heading 2')
 para(s['chapter']+'  /  '+s['placement'],'Caption')
 for txt in s['script']:para(txt)
 p=para('추가 질문  '+s['question']);p.paragraph_format.space_after=Pt(4)
 p=para('연결  '+s['next_connection'],'Caption');p.paragraph_format.space_after=Pt(4)
 p=para('출처  ','Caption')
 for j,id in enumerate(s['sources']):
  if j:p.add_run(' · ')
  hyperlink(p,id,REF[id].get('rechecked_url',REF[id]['url']))
 p.paragraph_format.space_after=Pt(15)
doc.save(DEST/'outputs/W8_outline_script_draft_v1.docx')
print('outline/page plan/script and DOCX written; sources',len(sources),'used',sum(bool(r['slides']) for r in sources))

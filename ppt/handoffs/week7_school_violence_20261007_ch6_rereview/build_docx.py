import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
DATA=json.loads((ROOT/'CONTENT.json').read_text())
import sys
OUT=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'W7_6장_화면과대본_전페이지재검수_v2.docx'
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.5);sec.bottom_margin=Cm(1.45)
sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Pretendard';st.font.color.rgb=RGBColor.from_string('202A39')
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Pretendard')
 st.paragraph_format.line_spacing=1.1;st.paragraph_format.space_after=Pt(5)
for name,size in [('Normal',10.5),('Heading 1',17),('Heading 2',11.5),('Heading 3',10.5),('Caption',9)]:
 doc.styles[name].font.size=Pt(size)
for name in ['Heading 1','Heading 2','Heading 3']:
 doc.styles[name].font.color.rgb=RGBColor.from_string('25365A')
 doc.styles[name].paragraph_format.keep_with_next=True
 doc.styles[name].paragraph_format.space_before=Pt(6)
 doc.styles[name].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.widow_control=True
hp=sec.header.paragraphs[0];hp.text='7주차 학교폭력   ·   6장 처벌을 달리할 기준과 학생의 권리'
hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor.from_string('5A6475')
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('PPT 구성안 · 2026.10.07  |  ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)

def line(text, label=None, size=None, color=None, bold=False):
 p=doc.add_paragraph()
 if label:p.add_run(label+'  ').bold=True
 r=p.add_run(text);r.bold=bold
 if size:r.font.size=Pt(size)
 if color:r.font.color.rgb=RGBColor.from_string(color)
 return p

def table(sl):
 t=doc.add_table(rows=1, cols=len(sl['headers']));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 widths=[17.6*w/sum(sl['widths']) for w in sl['widths']]
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 allrows=[sl['headers']]+sl['rows']
 for i,vals in enumerate(allrows):
  row=t.rows[0] if i==0 else t.add_row()
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for cell,w,txt in zip(row.cells,widths,vals):
   cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'25365A' if i==0 else ('F3F5F8' if i%2==0 else 'FFFFFF'));cp.append(sh)
   mar=OxmlElement('w:tcMar')
   for side,val in [('top','90'),('bottom','90'),('left','105'),('right','105')]:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');mar.append(el)
   cp.append(mar);cell.text=txt
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.07
    for r in p.runs:
     r.font.size=Pt(9.7)
     if i==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
     elif len(vals)==2 and sl['id']!='6-05' and vals.index(txt)==0:r.bold=True
  row.height=None
 borders=OxmlElement('w:tblBorders')
 for side in ('top','bottom','left','right','insideH','insideV'):
  b=OxmlElement('w:'+side)
  for k,v in [('val','single'),('sz','4'),('color','D4DCE5')]:b.set(qn('w:'+k),v)
  borders.append(b)
 t._tbl.tblPr.append(borders)

for i,sl in enumerate(DATA['slides']):
 if i:doc.add_page_break()
 doc.add_heading(f'{sl["id"]}  {sl["title"]}',1)
 doc.add_heading('PPT 구성안',2)
 # Layout instructions preserved in CONTENT.json, omitted from the reading copy
 if sl['band']:
  p=line(sl['band'],'하단 질문' if sl['id']=='6-10' else '상단 문장',bold=True)
 if sl['context']:line(sl['context'],'본문 맥락',size=9.5)
 table(sl)
 for b in sl['below']:line(b,size=9.7)
 doc.add_heading('낭독 대본',2)
 for txt in sl['script']:line(txt)
doc.core_properties.title='7주차 6장 PPT 구성안과 낭독 대본'
doc.core_properties.subject='처벌을 달리할 기준과 학생의 권리 · 10개 화면'
doc.core_properties.author=''
doc.core_properties.comments='화면·대본은 2026-10-07 검토본. 출처·상세자료·검수기록은 별도 인수인계 묶음에 보존. 실제 PPTX 아님.'
OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
print(OUT)

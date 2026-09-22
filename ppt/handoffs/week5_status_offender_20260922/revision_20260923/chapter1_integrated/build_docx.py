from pathlib import Path
import json
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).parent
D=json.loads((BASE/'content.json').read_text())
doc=Document(); sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(19);sec.bottom_margin=Mm(18)
sec.left_margin=sec.right_margin=Mm(21)
sec.header_distance=sec.footer_distance=Mm(8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Pretendard';st.font.color.rgb=RGBColor(0,0,0)
 rf=st.element.get_or_add_rPr().rFonts
 for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Pretendard')
 for key in list(rf.attrib):
  if 'theme' in key.lower():del rf.attrib[key]
 st.paragraph_format.line_spacing=1.18;st.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(10.5)
for name,size in [('Title',24),('Heading 1',18),('Heading 2',11.5),('Heading 3',10.5)]:
 doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(4)
hdr=sec.header.paragraphs[0];hdr.text='경찰대학  |  범죄·공공안전세미나Ⅱ 5주차'
hdr.runs[0].font.size=Pt(8)
ft=sec.footer.paragraphs[0];ft.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=ft.add_run('제1장 통합 보강안 · 강의 대본  |  ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');ft._p.append(f)

def p(text,bold=False,size=None,after=7):
 para=doc.add_paragraph();para.paragraph_format.space_after=Pt(after)
 r=para.add_run(text);r.bold=bold
 if size:r.font.size=Pt(size)
 return para
def h(t):doc.add_heading(t,2)
def newpage(t):doc.add_page_break();doc.add_heading(t,1)
def link(label,url):
 para=doc.add_paragraph();para.paragraph_format.space_after=Pt(6)
 rel=doc.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rel)
 rr=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 font=OxmlElement('w:rFonts')
 for a in ['ascii','hAnsi','eastAsia','cs']:font.set(qn('w:'+a),'Pretendard')
 pr.append(font);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz)
 u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
 rr.append(pr);tt=OxmlElement('w:t');tt.text=label;rr.append(tt);el.append(rr);para._p.append(el)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Mm(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  c=t.add_row().cells
  for i,x in enumerate(row):c[i].text=str(x)
 borders=OxmlElement('w:tblBorders')
 for k in ['top','bottom','insideH']:
  e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D0D4DC');borders.append(e)
 t._tbl.tblPr.append(borders)
 for j,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if j==0:pr.append(OxmlElement('w:tblHeader'))
  for i,c in enumerate(row.cells):
   c.width=Mm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    a=OxmlElement('w:'+edge);a.set(qn('w:w'),'80');a.set(qn('w:type'),'dxa');mar.append(a)
   cp.append(mar)
   if j==0:sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F4F8');cp.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.12
    for run in para.runs:run.font.size=Pt(9);run.bold=(j==0)
 return t

doc.add_paragraph('제1장 통합 PPT\n자료·강의 대본',style='Title')
p('논제의 대상과 현행 우범소년 제도',True,14)
p(D['proposition'],True,11)
p('11장 · 설명 약 24분 + 학생 발언 시간\n'+D['version'],size=9)
p('제도 이해 → 위험의 구별 → 절차 → 현장의 어려움 → 실제 사례 → 논거와 판단의 순서로 구성했다. 각 페이지의 대본은 바로 읽어 설명할 수 있도록 작성했으며, 원자료의 사실과 강의에서 던지는 질문을 구별했다.')
table(['PPT','내용','이후 논의에 남기는 질문'],[
('1–3','세 출발점·세 요건·법정 우범사유','가출 이유와 장래 우려를 무엇으로 판단하는가?'),
('4–5','두 위험의 구별·송치와 통고','피해 보호의 필요가 어떻게 소년사법 판단으로 이어지는가?'),
('6','반복 피해와 지원 연결의 어려움','개입을 지속하기 어렵게 만드는 조건은 무엇인가?'),
('7–8','시설 위탁 사례·7호 처분 사례','송치의 근거와 처분 선택의 이유는 각각 무엇인가?'),
('9','기관 협의로 달라진 송치 계획','피해지원 경로에서 누가 보호를 책임지는가?'),
('10–11','현장 논거·법 적용과 제도 판단','보호의 이익·자유 제한·다른 지원 가능성을 어떻게 비교하는가?')],[16,67,85])
h('강의에서 사용할 때')
p('PPT 화면에는 법조문과 자료의 성격 등 이해에 필요한 정보만 남겼다. 원자료의 위치·링크와 상세 설명은 이 문서 및 PPT 발표자 노트에 있다. 각 대본 끝의 연결 문장은 다음 화면으로 넘어갈 때 사용한다.',size=10)
p('실무자료와 토론회 사례는 법원의 상세 요건 판단이나 처분 효과까지 입증하지 않는다. 확인되지 않은 동기·진단·성과를 덧붙이지 않는다. 송치를 하지 않는 선택도 피해지원·주거·치료·가해자 수사 등 구체적인 대응과 함께 비교한다.',size=10)
p('확인 기준: master '+D['base_commit']+'\n전체 8개 단원은 유지. 제1장의 삭제된 논제 안/밖 페이지는 복원하지 않았다. 최신 두 사례 보강본의 내용과 배치를 이어받았다.',size=8.5)

for d in D['slides']:
 newpage(f"{d['n']:02d}  {d['title']}")
 p(f"PPT {d['n']}  |  {d['time']}  |  근거 "+', '.join(d['sources']),size=9)
 p(d['purpose'],True,10.5)
 h('강의 대본')
 for text in d['script']:p(text,size=11)
 h('자료에서 확인할 부분')
 p(d['evidence'],size=9.5)
 p('설명 범위  '+d['guard'],size=9.5)
 if d.get('optional'):
  p(d['optional'],size=9.5)
 h('다음 화면으로')
 p(d['transition'],True,10)
 src=[x for x in D['sources'] if x['id'] in d['sources']]
 # Source links are compact; bibliography gives exact page locations.
 for a in src:link('['+a['id']+'] '+a['name'],a['url'])

for idx,group in enumerate([D['sources'][:5],D['sources'][5:]]):
 newpage('자료 목록과 확인 위치 '+str(idx+1))
 for a in group:
  h('['+a['id']+'] '+a['name'])
  p(a['where'],size=10)
  p('활용  '+a['use'],size=10)
  link('원자료 확인',a['url'])
  if a.get('extra'):p(a['extra'],size=8.5)
 if idx==1:
  h('뒤 단원으로 넘길 자료와 질문')
  table(['단원','제1장 자료의 재사용 범위'],[
   ('제2장','피해청소년의 법적 지위 변화와 우범소년 경로의 관계를 법 개정 내용으로 설명한다.'),
   ('제3·4장','피해 규모 원표와 현장 맥락, 송치·처분의 실제 분포를 제시한다. 사례를 대표 통계로 바꾸지 않는다.'),
   ('제5·6장','숙소·치료·생활관리·자유 제한을 비교하고, 연결 시점·인력·주거·기관 책임을 구체화한다.'),
   ('제7·8장','효과 연구의 대상과 결과를 확인한 뒤 논거를 평가한다. 사례의 처리 결과를 효과 증거로 대체하지 않는다.')],[25,143])

newpage('법조문 확인 · 소년법 제4조')
p('제4조(보호의 대상과 송치 및 통고)',True,12)
law=[
'① 다음 각 호의 어느 하나에 해당하는 소년은 소년부의 보호사건으로 심리한다.',
'1. 죄를 범한 소년',
'2. 형벌 법령에 저촉되는 행위를 한 10세 이상 14세 미만인 소년',
'3. 다음 각 목에 해당하는 사유가 있고 그의 성격이나 환경에 비추어 앞으로 형벌 법령에 저촉되는 행위를 할 우려가 있는 10세 이상인 소년',
'가. 집단적으로 몰려다니며 주위 사람들에게 불안감을 조성하는 성벽(性癖)이 있는 것',
'나. 정당한 이유 없이 가출하는 것',
'다. 술을 마시고 소란을 피우거나 유해환경에 접하는 성벽이 있는 것',
'② 제1항제2호 및 제3호에 해당하는 소년이 있을 때에는 경찰서장은 직접 관할 소년부에 송치(送致)하여야 한다.',
'③ 제1항 각 호의 어느 하나에 해당하는 소년을 발견한 보호자 또는 학교ㆍ사회복리시설ㆍ보호관찰소(보호관찰지소를 포함한다. 이하 같다)의 장은 이를 관할 소년부에 통고할 수 있다.',
'[전문개정 2007. 12. 21.]']
for t in law:p(t,size=10.5,after=7)
link('국가법령정보센터 · 제4조 원문','https://www.law.go.kr/lsLawLinkInfo.do?chrClsCd=010202&lsJoLnkSeq=900434284')
h('두 사례의 처분 명칭을 읽는 기준')
p('제32조 제1항 제5호는 장기 보호관찰, 제6호는 아동복지시설이나 그 밖의 소년보호시설에 감호 위탁, 제7호는 병원·요양소 또는 의료재활소년원 위탁이다. 이 설명은 각 호를 간추린 것이며, 아래 원문에서 정확한 법정 명칭과 병합 가능 처분을 확인할 수 있다.',size=10)
link('국가법령정보센터 · 제32조 원문',D['sources'][1]['url'])
p('형사처벌과 보호처분을 혼동하지 않는다. 분류심사원 임시위탁은 조사·심리 과정의 조치이며 최종 처분 번호와 구분한다. 7호를 ‘치료감호’라는 정식 처분명으로 표시하지 않는다.',size=10)

# Explicit font families on runs and linked character styles prevent renderer fallback.
for st in doc.styles:
 try:
  st.font.name='Pretendard'
  rf=st.element.get_or_add_rPr().rFonts
  for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Pretendard')
  for key in list(rf.attrib):
   if 'theme' in key.lower():del rf.attrib[key]
 except AttributeError:pass
for root in [doc.element,sec.header._element,sec.footer._element]:
 for run in root.iter(qn('w:r')):
  rp=run.find(qn('w:rPr'))
  if rp is None:rp=OxmlElement('w:rPr');run.insert(0,rp)
  rf=rp.find(qn('w:rFonts'))
  if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
  for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Pretendard')
  for key in list(rf.attrib):
   if 'theme' in key.lower():del rf.attrib[key]
out=BASE/'generated/W5_CH1_통합보강_자료와대본_v1.docx';out.parent.mkdir(exist_ok=True);doc.save(out)
print(out)

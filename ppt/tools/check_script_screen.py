#!/usr/bin/env python3
"""화면-대본 같은 말 대조 도구 (PPT_CONTENT_RULES §4-2-2 보조).

사용법:
    python3 check_script_screen.py DECK.pptx [--pages 3-8,12] [--min 0.6] [--out REPORT.md]

검사 내용
  1. 대본 문단마다 핵심어(2글자 이상 낱말에서 조사·어미를 뗀 줄기)가
     같은 장 화면에 같은 말로 있는지 비율을 계산하고, 화면에 없는 낱말을 나열한다.
  2. 대본이 화면 위치(왼쪽·오른쪽·위·아래·띠·칸·번호)를 가리키는 문단을 표시한다.
  3. 화면·대본의 방어문(해석 경계문) 후보를 찾는다.

이 도구는 낱말 대조와 문형 탐지만 한다. 의미가 같은지, 위치 지시가 실제
자리와 맞는지, 탐지된 부정문이 자료 자체의 내용(법적 금지·권리·제외 조항)인지는
사람이 화면을 읽고 판정한다. 결과를 의미 검수 PASS의 근거로 쓰지 않는다.

발표자 노트에 '[낭독 대본]' 표시가 있으면 그 구역(다음 '[' 머리표 전까지)만
대본으로 읽고, 없으면 노트 전체를 대본으로 읽는다.
"""
import argparse, re, sys
from pptx import Presentation

JOSA = sorted("""은 는 이 가 을 를 의 에 에서 에게 에게서 께 으로 로 와 과 도 만 까지 부터 이나 나
이라 라 이며 며 이고 고 보다 처럼 마다 조차 이다 입니다 였다 이었다 이었습니다 였습니다 합니다 한다 했다
했습니다 됩니다 된다 되었다 되었습니다 됐다 습니다 ㅂ니다 니다 해서 하고 하면 하는 하여 해 할 한 된 될
되는 되어 돼 인 일 임 이란 란 라는 이라는 으로써 로써 으로서 로서 에는 에도 에서는 에서도 으로는 로는
와는 과는 이는 들 들은 들이 들을 들의""".split(), key=len, reverse=True)
STOP = set("""그리고 그러나 하지만 따라서 그래서 또한 다만 먼저 다음 이제 여기 여기서 지금 이것 그것 이런 그런
저런 이번 앞에서 앞의 뒤의 우리 여러분 있습니다 있다 없다 없습니다 있는 없는 같은 다른 경우 때문 대한
대해 위해 통해 함께 모두 가장 바로 보겠습니다 봅니다 살펴보겠습니다 보면 보았습니다 말합니다 뜻합니다
화면 왼쪽 오른쪽 아래 가운데 하나 두 세 네 것 수 등 및 또는 즉 곧 각각 이렇게 그렇게 어떻게 무엇
누가 어떤 얼마나 왜 언제 쪽 질문 정리 확인 번째 첫째 둘째 셋째 넷째 예를 들어 실제 현재 그중 칸 항목 문단 정도 때문에 때에 결과 내용 문제""".split())
POS = re.compile(r'(왼쪽|오른쪽|가운데|맨 위|위쪽|아래쪽|아래 띠|위 띠|띠|칸|첫째 줄|둘째 줄|[0-9]번|첫 번째|두 번째|세 번째)')
DEF = re.compile(r'(다루지 않|의미하지 않|뜻하지 않|뜻은 아니|뜻이 아니|것은 아니|것이 아니|인 것은 아|'
                 r'로 읽지 않|으로 읽지 않|라고 읽지|로 볼 수 없|으로 볼 수 없|해석하면 안|해석해서는|'
                 r'곧바로 대신하지|대신하지 않|과정으로 읽|만으로[^.]{0,40}(아니|않)|'
                 r'(?<!뿐 )(?<!뿐)[이가] 아니라|이 아닌 |가 아닌 )')


def stem(w):
    for j in JOSA:
        if len(w) > len(j) + 1 and w.endswith(j):
            return w[:-len(j)]
    return w


VERB = re.compile(r'(니다|습니까|겠|죠|요|다|고|며|면|서|지|게|도록|는데|므로|어야|아야|여야|려면|려고|으면|'
                  r'는지|은지|던|었|았|였|했|합|됩|됐|된|될|되|한|할|하|해|인지|든지|거나|더라도|지만|자|듯)$')


def words(t):
    out = []
    for w in re.findall(r'[가-힣A-Za-z0-9·]+', t):
        if re.fullmatch(r'[가-힣]+', w) and VERB.search(w) and not re.search(r'(법|조|항|호|권|자|센터|정보|기관|경찰|시민|사람|명단|기록|근거)$', w):
            continue
        s = stem(w)
        if len(s) >= 2 and s not in STOP and not s.isdigit():
            out.append(s)
    return out


def script_of(notes):
    if '[낭독 대본]' in notes:
        notes = notes.split('[낭독 대본]', 1)[1]
        notes = re.split(r'\n\[[^\]]+\]', notes)[0]
    return [p.strip() for p in re.split(r'\n\s*\n|\n', notes) if p.strip()]


def screen_of(slide):
    parts = []
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:  # group
                walk(sh.shapes)
            elif sh.has_text_frame:
                parts.append(sh.text_frame.text)
            elif getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        parts.append(c.text)
    walk(slide.shapes)
    return '\n'.join(parts)


def pages(spec, n):
    if not spec:
        return list(range(1, n + 1))
    out = []
    for part in spec.split(','):
        a, _, b = part.partition('-')
        out += list(range(int(a), int(b or a) + 1))
    return [p for p in out if 1 <= p <= n]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('deck'); ap.add_argument('--pages'); ap.add_argument('--min', type=float, default=0.6)
    ap.add_argument('--out')
    a = ap.parse_args()
    prs = Presentation(a.deck)
    sl = list(prs.slides)
    lines = [f'# 화면-대본 대조 보고: {a.deck}', '',
             f'기준: 대본 문단 핵심어의 화면 표시 비율 {a.min:.0%} 미만이면 [낮음]. 의미 판정은 사람이 한다.', '']
    low = dfn = 0
    for p in pages(a.pages, len(sl)):
        s = sl[p - 1]
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
        scr = screen_of(s)
        flat = re.sub(r'\s+', '', scr)
        paras = script_of(notes)
        lines.append(f'## {p}쪽')
        for i, para in enumerate(paras, 1):
            ws = list(dict.fromkeys(words(para)))
            if not ws:
                continue
            miss = [w for w in ws if w not in flat]
            cov = 1 - len(miss) / len(ws)
            tag = '[낮음] ' if cov < a.min else ''
            low += bool(tag)
            pos = ' [위치 지시: ' + ', '.join(sorted(set(POS.findall(para)))) + ']' if POS.search(para) else ''
            lines.append(f'- {tag}문단 {i}: {cov:.0%}{pos} — 화면에 없는 말: {", ".join(miss[:14]) or "없음"}')
        for src, txt in (('화면', scr), ('대본', '\n'.join(paras))):
            for sent in re.split(r'(?<=[.다])\s+|\n', txt):
                if DEF.search(sent):
                    dfn += 1
                    lines.append(f'- [방어문 후보·{src}] {sent.strip()[:160]}')
        lines.append('')
    lines.insert(3, f'요약: 낮은 문단 {low}개, 방어문 후보 {dfn}개.')
    rep = '\n'.join(lines)
    if a.out:
        open(a.out, 'w', encoding='utf-8').write(rep + '\n')
    print(rep if not a.out else f'낮은 문단 {low}개, 방어문 후보 {dfn}개 → {a.out}')


if __name__ == '__main__':
    main()
